# pyright: reportPrivateUsage=false
# ColorFormat return values are only partially typed; ignore those member diagnostics below.

"""Non-destructive font color queries and their serialization behavior."""

from __future__ import annotations

from io import BytesIO
from typing import TYPE_CHECKING, cast
from zipfile import ZipFile

import pytest

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL, MSO_THEME_COLOR
from pptx.shapes.autoshape import Shape
from pptx.text.text import Font
from pptx.util import Inches

from ..unitutil.cxml import element, xml

if TYPE_CHECKING:
    from pptx.oxml.text import CT_TextCharacterProperties
    from pptx.presentation import Presentation as PresentationType


class DescribeFontHasColor:
    @pytest.mark.parametrize("rPr_tag", ["a:rPr", "a:defRPr", "a:endParaRPr"])
    @pytest.mark.parametrize(
        ("fill_cxml", "expected"),
        [
            ("", False),
            ("/a:noFill", False),
            ("/a:gradFill/a:gsLst/a:gs{pos=0}/a:srgbClr{val=123456}", False),
            ("/a:pattFill/a:fgClr/a:srgbClr{val=123456}", False),
            ("/a:blipFill", False),
            ("/a:grpFill", False),
            ("/a:highlight/a:srgbClr{val=123456}", False),
            ("/a:uFill/a:solidFill/a:srgbClr{val=123456}", False),
            ("/a:solidFill", True),
            ("/a:solidFill/a:srgbClr{val=123456}", True),
            ("/a:solidFill/a:schemeClr{val=accent1}/a:lumMod{val=75000}", True),
        ],
    )
    def it_queries_solid_fill_presence_without_changing_xml(
        self, rPr_tag: str, fill_cxml: str, expected: bool
    ):
        rPr = cast("CT_TextCharacterProperties", element(rPr_tag + fill_cxml))
        font = Font(rPr)
        original_xml = rPr.xml

        for _ in range(3):
            assert font.has_color is expected
            assert rPr.xml == original_xml

    @pytest.mark.parametrize(
        ("color_cxml", "expected_type"),
        [
            ("", None),
            ("/a:srgbClr{val=123456}", MSO_COLOR_TYPE.RGB),
            ("/a:schemeClr{val=accent1}/a:lumMod{val=75000}", MSO_COLOR_TYPE.SCHEME),
            ("/a:scrgbClr{r=10000,g=20000,b=30000}", MSO_COLOR_TYPE.SCRGB),
            ("/a:hslClr{hue=1200000,sat=50000,lum=50000}", MSO_COLOR_TYPE.HSL),
            ("/a:sysClr{val=windowText,lastClr=000000}", MSO_COLOR_TYPE.SYSTEM),
            ("/a:prstClr{val=red}", MSO_COLOR_TYPE.PRESET),
        ],
    )
    def it_allows_guarded_color_reads_without_changing_xml(
        self, color_cxml: str, expected_type: MSO_COLOR_TYPE | None
    ):
        rPr = cast("CT_TextCharacterProperties", element("a:rPr/a:solidFill" + color_cxml))
        font = Font(rPr)
        original_xml = rPr.xml

        for _ in range(3):
            assert font.has_color is True
            assert font.color.type == expected_type  # pyright: ignore[reportUnknownMemberType]
            assert rPr.xml == original_xml

    def it_reflects_fill_changes_without_caching_the_answer(self):
        rPr = cast("CT_TextCharacterProperties", element("a:rPr"))
        font = Font(rPr)
        assert font.has_color is False

        font.fill.solid()
        assert font.has_color is True
        assert font.color.type is None  # pyright: ignore[reportUnknownMemberType]

        font.fill.background()
        assert font.has_color is False

        font.fill.gradient()
        assert font.has_color is False

        # Another Font can reference the same properties element. Even when this
        # font has a cached FillFormat, the query must reflect the current XML.
        other_font = Font(rPr)
        other_font.color.rgb = RGBColor(0x12, 0x34, 0x56)
        assert font.has_color is True

        other_font.fill.patterned()
        assert font.has_color is False

    @pytest.mark.parametrize("fill_cxml", ["", "/a:noFill", "/a:gradFill", "/a:pattFill"])
    @pytest.mark.parametrize("color_type", ["rgb", "theme"])
    def it_preserves_chained_color_setters(self, fill_cxml: str, color_type: str):
        font = Font(cast("CT_TextCharacterProperties", element("a:rPr" + fill_cxml)))
        assert font.has_color is False

        if color_type == "rgb":
            font.color.rgb = RGBColor(0x12, 0x34, 0x56)
            expected_cxml = "a:rPr/a:solidFill/a:srgbClr{val=123456}"
            assert font.color.type == MSO_COLOR_TYPE.RGB  # pyright: ignore[reportUnknownMemberType]
            assert font.color.rgb == RGBColor(0x12, 0x34, 0x56)
        else:
            font.color.theme_color = MSO_THEME_COLOR.ACCENT_1
            expected_cxml = "a:rPr/a:solidFill/a:schemeClr{val=accent1}"
            expected_type = MSO_COLOR_TYPE.SCHEME
            assert font.color.type == expected_type  # pyright: ignore[reportUnknownMemberType]
            assert font.color.theme_color == MSO_THEME_COLOR.ACCENT_1

        assert font.has_color is True
        assert font.fill.type == MSO_FILL.SOLID
        assert font._rPr.xml == xml(expected_cxml)

    @pytest.mark.parametrize("inherited_color_type", ["rgb", "theme"])
    def it_preserves_inherited_color_through_save_and_reload(self, inherited_color_type: str):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        paragraph = textbox.text_frame.paragraphs[0]
        if inherited_color_type == "rgb":
            paragraph.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        else:
            paragraph.font.color.theme_color = MSO_THEME_COLOR.ACCENT_1
        run = paragraph.add_run()
        run.text = "Inherited color"
        assert run._r.rPr is None

        # Acquiring Font can add an empty rPr; has_color must not add a fill.
        font = run.font
        original_xml = paragraph._p.xml
        original_slide_xml = self._slide_xml(prs)
        for _ in range(3):
            assert font.has_color is False
            assert (font.color if font.has_color else None) is None
            assert paragraph._p.xml == original_xml
        assert self._slide_xml(prs) == original_slide_xml

        stream = BytesIO()
        prs.save(stream)
        reloaded = Presentation(stream)
        shape = cast(Shape, reloaded.slides[0].shapes[0])
        paragraph = shape.text_frame.paragraphs[0]
        font = paragraph.runs[0].font
        assert font.has_color is False
        assert font.fill.type is None
        assert paragraph._p.xml == original_xml
        color = paragraph.font.color
        if inherited_color_type == "rgb":
            expected_rgb = RGBColor(0xFF, 0xFF, 0xFF)
            assert color.rgb == expected_rgb  # pyright: ignore[reportUnknownMemberType]
        else:
            expected_theme = MSO_THEME_COLOR.ACCENT_1
            assert color.theme_color == expected_theme  # pyright: ignore[reportUnknownMemberType]

        # A later change to the paragraph color must still leave the run with no
        # direct fill override, including after another save/reload cycle.
        paragraph.font.color.rgb = RGBColor(0x12, 0x34, 0x56)
        stream = BytesIO()
        reloaded.save(stream)
        shape = cast(Shape, Presentation(stream).slides[0].shapes[0])
        paragraph = shape.text_frame.paragraphs[0]
        color = paragraph.font.color
        assert color.rgb == RGBColor(0x12, 0x34, 0x56)  # pyright: ignore[reportUnknownMemberType]
        assert paragraph.runs[0].font.has_color is False
        assert paragraph.runs[0].font.fill.type is None

    @pytest.mark.parametrize(
        ("fill_cxml", "expected"),
        [
            ("a:solidFill", True),
            ("a:solidFill/a:srgbClr{val=123456}", True),
            ("a:solidFill/a:schemeClr{val=accent1}/a:lumMod{val=75000}", True),
            ("a:noFill", False),
            ("a:gradFill/a:gsLst/a:gs{pos=0}/a:srgbClr{val=123456}", False),
            ("a:pattFill/a:fgClr/a:srgbClr{val=123456}", False),
        ],
    )
    def it_preserves_direct_fill_through_save_and_reload(self, fill_cxml: str, expected: bool):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        run = shape.text_frame.paragraphs[0].add_run()
        run.text = "Direct fill"
        font = run.font
        font._rPr.append(element(fill_cxml))
        original_xml = run._r.xml
        original_slide_xml = self._slide_xml(prs)

        for _ in range(3):
            assert font.has_color is expected
            if font.has_color:
                _ = font.color
        assert run._r.xml == original_xml
        assert self._slide_xml(prs) == original_slide_xml

        stream = BytesIO()
        prs.save(stream)
        reloaded = Presentation(stream)
        shape = cast(Shape, reloaded.slides[0].shapes[0])
        run = shape.text_frame.paragraphs[0].runs[0]
        assert run.font.has_color is expected
        assert run._r.xml == original_xml

    @staticmethod
    def _slide_xml(prs: PresentationType) -> bytes:
        stream = BytesIO()
        prs.save(stream)
        with ZipFile(stream) as package:
            return package.read("ppt/slides/slide1.xml")
